"""Catalogue imports (POC_05 §3.2–3.4): file → extracted items → validated → matched → draft change set → summary.

The Temporal `ImportWorkflow` (workflows module) drives these steps as activities; everything here is idempotent per
`import_run_id` so a retried activity does not duplicate rows.

* Spreadsheets/CSV: deterministic column detection by header names; the extract role is asked for a column map only
  when the header is unclear.
* Images/PDFs: the vision role, schema-constrained, told to copy prices exactly and never guess.
* Every price must be readable from the item's own source text (`price_not_in_source` catches hallucination).
"""

from __future__ import annotations

import csv
import io
import json
import re
from typing import Any, Literal
from uuid import UUID

from pydantic import BaseModel, Field
from sqlalchemy import text

from reflexo.common import ports
from reflexo.common.contracts import ModelGatewayPort
from reflexo.common.db import get_db
from reflexo.common.ids import new_id
from reflexo.common.llm import CallCtx, ChatMessage, MediaInput
from reflexo.common.logging import get_logger
from reflexo.common.money import format_inr
from reflexo.knowledge import changesets
from reflexo.knowledge.query import canonical_size
from reflexo.knowledge.resolve import resolve_offerings

log = get_logger(__name__)
MAX_PRICE_PAISE = 10_00_000 * 100  # ₹10,00,000
MATCH_THRESHOLD = 0.6
_DEVANAGARI = str.maketrans("०१२३४५६७८९", "0123456789")


class ExtractedVariant(BaseModel):
    name: str = "default"  # "1 kg", "box of 6", "default"
    price_text: str | None = None  # exactly as printed
    price_paise: int | None = None  # filled by validation


class ExtractedItem(BaseModel):
    name: str = Field(min_length=1, max_length=200)
    category: str | None = None
    description: str | None = None
    variants: list[ExtractedVariant] = Field(default_factory=list)
    price_on_request: bool = False
    source_text: str = ""
    confidence: float = Field(default=0.9, ge=0, le=1)


class Extraction(BaseModel):
    items: list[ExtractedItem]


class ColumnMap(BaseModel):
    name_col: int
    price_col: int
    variant_col: int | None = None
    category_col: int | None = None
    description_col: int | None = None
    confidence: float = Field(ge=0, le=1)


class ImportFailed(Exception):
    """A file we cannot read; the message is owner-facing."""


Status = Literal["new", "changed", "unchanged", "needs_input"]

# ---------------------------------------------------------------- prices


def normalize_price(raw: str | float | int | None) -> tuple[int | None, str | None]:
    """(paise, issue). Accepts ₹1,499 · Rs. 1499/- · 1499 rs · 1.5k · 1,10,000 · Devanagari digits.
    A range (500–700) or anything unreadable gives (None, issue)."""
    if raw is None or (isinstance(raw, str) and not raw.strip()):
        return None, "missing_price"
    if isinstance(raw, int | float):
        value = float(raw)
    else:
        s = raw.translate(_DEVANAGARI).lower().replace("₹", " ").replace("rs.", " ").replace("rs", " ")
        s = s.replace("inr", " ").replace("/-", " ").replace("rupees", " ").replace("rupaye", " ").strip()
        if re.search(r"\d\s*[-–—]\s*\d", s) or " to " in s:
            return None, "price_range"
        m = re.search(r"(\d[\d,]*(?:\.\d+)?)\s*(k)?\b", s)
        if not m:
            return None, "missing_price"
        value = float(m.group(1).replace(",", "")) * (1000 if m.group(2) else 1)
    paise = round(value * 100)
    if paise <= 0 or paise > MAX_PRICE_PAISE:
        return None, "price_out_of_range"
    return paise, None


def _digits(s: str) -> str:
    return re.sub(r"\D", "", s.translate(_DEVANAGARI))


# ---------------------------------------------------------------- extraction


_HEADER_WORDS: dict[str, tuple[str, ...]] = {
    "name": ("item", "name", "product", "particular", "dish", "cake", "menu"),
    "price": ("price", "rate", "mrp", "amount", "cost", "rs", "₹", "inr"),
    "variant": ("size", "weight", "variant", "qty", "quantity", "pack", "unit", "portion"),
    "category": ("category", "type", "section", "group"),
    "description": ("description", "desc", "details", "notes"),
}


def _sheet_rows(content: bytes, mime: str, filename: str | None) -> list[list[str]]:
    name = (filename or "").lower()
    if mime.startswith("text/csv") or name.endswith(".csv"):
        return [
            [c.strip() for c in r]
            for r in csv.reader(io.StringIO(content.decode("utf-8-sig", errors="replace")))
            if any(c.strip() for c in r)
        ]
    if name.endswith(".xls") and not name.endswith(".xlsx"):
        raise ImportFailed("This is an old .xls file — please send it as .xlsx or CSV.")
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ImportFailed("I couldn't open this Excel file — is it password-protected?") from exc
    for ws in wb.worksheets:  # first sheet with ≥ 2 non-empty columns
        rows = [
            ["" if v is None else str(v).strip() for v in row]
            for row in ws.iter_rows(values_only=True)
            if any(v not in (None, "") for v in row)
        ]
        if rows and max(sum(1 for c in r if c) for r in rows) >= 2:
            wb.close()
            return rows
    wb.close()
    raise ImportFailed("The sheet looks empty.")


def _guess_columns(header: list[str]) -> ColumnMap | None:
    found: dict[str, int] = {}
    for i, cell in enumerate(header):
        low = cell.lower()
        for role, words in _HEADER_WORDS.items():
            if role not in found and any(w in low for w in words):
                found[role] = i
                break
    if "name" not in found or "price" not in found:
        return None
    return ColumnMap(
        name_col=found["name"],
        price_col=found["price"],
        variant_col=found.get("variant"),
        category_col=found.get("category"),
        description_col=found.get("description"),
        confidence=0.95,
    )


async def _llm_columns(business_id: UUID, header: list[str], sample: list[list[str]]) -> ColumnMap:
    prompt = (
        "A shop's price list spreadsheet. Return the 0-based column index of the item name and the price, and of the "
        "size/variant, category and description if present (else null), plus your confidence 0..1.\n"
        f"Header: {json.dumps(header, ensure_ascii=False)}\nRows: {json.dumps(sample, ensure_ascii=False)}"
    )
    model = ports.registry.require(ModelGatewayPort)  # type: ignore[type-abstract]
    return await model.structured(
        "extract",
        [ChatMessage(role="user", content=prompt)],
        ColumnMap,
        ctx=CallCtx(business_id=business_id, purpose="import.columns"),
    )


async def extract_spreadsheet(
    business_id: UUID, content: bytes, mime: str, filename: str | None
) -> list[ExtractedItem]:
    rows = _sheet_rows(content, mime, filename)
    header_idx = next((i for i, r in enumerate(rows) if sum(1 for c in r if c and not _is_number(c)) >= 2), 0)
    header, body = rows[header_idx], rows[header_idx + 1 :]
    cols = _guess_columns(header)
    if cols is None:
        cols = await _llm_columns(business_id, header, body[:5])
        if cols.confidence < 0.8:
            names = ", ".join(f"{i + 1}) {h}" for i, h in enumerate(header) if h)
            raise ImportFailed(f"Which column has the price? {names}")
    by_name: dict[str, ExtractedItem] = {}
    for r in body:
        name = _cell(r, cols.name_col)
        if not name:
            continue
        variant = ExtractedVariant(
            name=_cell(r, cols.variant_col) or "default", price_text=_cell(r, cols.price_col) or None
        )
        item = by_name.get(name.lower())
        if item is None:
            item = ExtractedItem(
                name=name,
                category=_cell(r, cols.category_col) or None,
                description=_cell(r, cols.description_col) or None,
                source_text=" | ".join(c for c in r if c),
                confidence=cols.confidence,
            )
            by_name[name.lower()] = item
        else:
            item.source_text += " || " + " | ".join(c for c in r if c)
        item.variants.append(variant)
    return list(by_name.values())


def _cell(row: list[str], i: int | None) -> str:
    return row[i].strip() if i is not None and i < len(row) else ""


def _is_number(s: str) -> bool:
    return normalize_price(s)[0] is not None and not re.search(r"[a-zA-Z]{3,}", s)


EXTRACT_PROMPT = (
    "Read this menu / price list and return every item. Copy prices EXACTLY as printed into price_text (with the "
    "currency sign if printed). Sizes or packs (500 g, 1 kg, box of 6) are separate variants of one item. If no price "
    "is printed for an item, set price_on_request=true and leave variants without price_text. Do NOT guess prices. "
    "source_text = the exact text you read for that item (name, size and price). category = the section heading if "
    "any. confidence 0..1 per item."
)


async def extract_visual(business_id: UUID, content: bytes, mime: str) -> list[ExtractedItem]:
    model = ports.registry.require(ModelGatewayPort)  # type: ignore[type-abstract]
    result = await model.vision(
        "vision",
        EXTRACT_PROMPT,
        [MediaInput(mime=mime, data=content)],
        Extraction,
        ctx=CallCtx(business_id=business_id, purpose="import.extract"),
    )
    return result.items


# ---------------------------------------------------------------- validate + match + draft


def validate(items: list[ExtractedItem], *, visual: bool) -> list[tuple[ExtractedItem, list[str]]]:
    """Normalize prices into paise; flag missing/unreadable prices, prices not in the source, duplicates."""
    seen: set[str] = set()
    out = []
    for item in items:
        issues: list[str] = []
        key = re.sub(r"\s+", " ", item.name.lower()).strip()
        if key in seen:
            issues.append("possible_duplicate")
        seen.add(key)
        if not item.variants:
            item.variants = [ExtractedVariant()]
        for v in item.variants:
            if v.price_text is None and v.price_paise is None:
                continue
            paise, issue = normalize_price(v.price_text)
            v.price_paise = paise
            if issue:
                issues.append(issue)
            elif visual and paise is not None and _digits(str(paise // 100)) not in _digits(item.source_text):
                issues.append("price_not_in_source")  # the model "read" a price that isn't in its own source text
        if item.price_on_request or all(v.price_paise is None for v in item.variants):
            issues.append("missing_price")
        out.append((item, sorted(set(issues))))
    return out


def _price_of(v: Any) -> int | None:
    return int(v) if v is not None else None


async def prepare(
    business_id: UUID,
    import_run_id: UUID,
    items: list[ExtractedItem],
    *,
    visual: bool,
    origin: str,
    member_id: UUID | None,
    filename: str | None,
) -> dict[str, Any]:
    """Validate, match against the catalogue, write `import_items`, build the draft change set.
    Returns {change_set_id|None, lines, needs_input, stats}. Idempotent per run (re-running rebuilds the draft)."""
    checked = validate(items, visual=visual)
    cs = await changesets.begin(business_id, origin, f"Import {filename or 'file'}", member_id=member_id)
    needs: list[str] = []
    stats = {"found": len(checked), "new": 0, "changed": 0, "unchanged": 0, "needs_review": 0}
    rows: list[dict[str, Any]] = []
    for item, issues in checked:
        blocking = [i for i in issues if i != "possible_duplicate"] or (["possible_duplicate"] if issues else [])
        status: Status
        target: UUID | None = None
        if blocking:
            status = "needs_input"
            reason = {
                "missing_price": "price not found",
                "price_range": "price is a range",
                "price_not_in_source": "price unclear",
                "price_out_of_range": "price looks wrong",
                "possible_duplicate": "listed twice",
            }[blocking[0]]
            needs.append(f"{item.name} — {reason}")
        else:
            status, target = await _apply_item(cs, business_id, item)
        stats["needs_review" if status == "needs_input" else status] += 1
        rows.append(
            {
                "id": new_id(),
                "extracted": item.model_dump_json(),
                "confidence": round(item.confidence, 2),
                "issues": issues,
                "target": target,
            }
        )
    async with get_db().tenant_session(business_id) as s:
        await s.execute(text("DELETE FROM import_items WHERE import_run_id = :r"), {"r": import_run_id})
        for r in rows:
            await s.execute(
                text("""INSERT INTO import_items (id, business_id, import_run_id, extracted, confidence, issues,
                                                  target_offering_id)
                        VALUES (:id, :b, :r, CAST(:e AS jsonb), :c, :i, :t)"""),
                {
                    "id": r["id"],
                    "b": business_id,
                    "r": import_run_id,
                    "e": r["extracted"],
                    "c": r["confidence"],
                    "i": r["issues"],
                    "t": r["target"],
                },
            )
        preview = await changesets.preview(business_id, cs.id)
        has_changes = bool(preview.lines)
        await s.execute(
            text("""UPDATE import_runs SET stats = CAST(:st AS jsonb), change_set_id = :cs WHERE id = :r"""),
            {"st": json.dumps(stats), "cs": cs.id if has_changes else None, "r": import_run_id},
        )
    return {
        "change_set_id": str(cs.id) if has_changes else None,
        "lines": preview.lines,
        "needs_input": needs,
        "stats": stats,
    }


async def _apply_item(cs: Any, business_id: UUID, item: ExtractedItem) -> tuple[Status, UUID | None]:
    """Add create/update items for one extracted item; returns its status and matched offering."""
    priced = [v for v in item.variants if v.price_paise is not None]
    candidates = await resolve_offerings(business_id, item.name)
    best = candidates[0] if candidates and candidates[0].score >= MATCH_THRESHOLD else None
    if best is None:
        single_default = len(priced) == 1 and priced[0].name.lower() in ("default", "", "regular")
        created = await changesets.add(
            cs,
            "offering",
            "create",
            None,
            {
                "name": item.name,
                "category": item.category or "other",
                "description": item.description or "",
                "base_price_paise": priced[0].price_paise if single_default else None,
            },
        )
        assert created is not None
        if not single_default:
            for v in priced:
                await changesets.add(
                    cs,
                    "variant",
                    "create",
                    None,
                    {"offering_id": str(created.entity_id), "name": v.name, "price_paise": v.price_paise},
                )
        return "new", None
    changed = False
    for v in priced:
        if best.variants:
            wanted = canonical_size(v.name) if v.name.lower() != "default" else None
            existing = next(
                (e for e in best.variants if wanted and canonical_size(e.name) == wanted),
                best.variants[0] if len(best.variants) == 1 and v.name.lower() == "default" else None,
            )
            if existing is None:
                await changesets.add(
                    cs,
                    "variant",
                    "create",
                    None,
                    {"offering_id": str(best.offering_id), "name": v.name, "price_paise": v.price_paise},
                )
                changed = True
            elif _price_of(existing.price_paise) != v.price_paise:
                await changesets.add(cs, "variant", "update", existing.variant_id, {"price_paise": v.price_paise})
                changed = True
        elif _price_of(best.base_price_paise) != v.price_paise:
            await changesets.add(cs, "offering", "update", best.offering_id, {"base_price_paise": v.price_paise})
            changed = True
    return ("changed" if changed else "unchanged"), best.offering_id


def summary_text(filename: str | None, prepared: dict[str, Any], code: str | None) -> str:
    lines = prepared["lines"]
    head = f"📄 *{filename or 'Your file'}*: " + (
        f"{len(lines)} change{'s' if len(lines) != 1 else ''}" if lines else "no changes"
    )
    body = [f"• {ln.lstrip('✏️➕🗑 ').strip()}" for ln in lines]
    if prepared["needs_input"]:
        body.append("⚠️ Needs your input (left out for now):")
        body += [f"• {n}" for n in prepared["needs_input"]]
    tail = f"Reply *YES {code}* to apply, or tell me what to change." if code else ""
    return "\n".join([head, *body, *([tail] if tail else [])])


__all__ = [
    "ExtractedItem",
    "ExtractedVariant",
    "ImportFailed",
    "extract_spreadsheet",
    "extract_visual",
    "format_inr",
    "normalize_price",
    "prepare",
    "summary_text",
    "validate",
]
