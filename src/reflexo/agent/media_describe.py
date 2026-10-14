"""`media.describe` job (POC_05 §2, POC_04 §8.2): one line per file, shared by owner files and customer photos.

* spreadsheets — deterministic: sheet, header row, row count, first rows ("Excel: 3 rows — Item, Price; …");
* images and PDFs — the vision role (owner files: what the file is and its key contents; customer photos: what the
  product looks like, for matching and for the agent's prompt);
* anything else — its kind.
The description lands on `media_objects` (prompts show `[image: …]`) and, for owner files, on `owner_files`.
"""

from __future__ import annotations

import asyncio
import csv
import io
from pathlib import Path
from uuid import UUID

from pydantic import BaseModel, Field

import reflexo.gateway.api as gateway
import reflexo.owner.api as owner
from reflexo.agent.model_gateway.gateway import ModelGateway
from reflexo.common.enums import PipelineEventStatus
from reflexo.common.llm import CallCtx, MediaInput, ModelError
from reflexo.common.logging import get_logger
from reflexo.common.xray import emit

log = get_logger(__name__)
QUEUE = "media.describe"
MAX_INLINE_BYTES = 15 * 1024 * 1024
OWNER_PROMPT = (
    "Describe this file for a shop owner's file library in ONE short line (max 25 words): what it is (menu, price "
    "list, product photo, receipt, invoice…) and its key contents (item names and prices if visible)."
)
CUSTOMER_PROMPT = (
    "A customer sent this photo to a shop. Describe it in ONE short line (max 30 words) for the shop assistant: "
    "what product it shows (type, flavour cues, colours, decorations, size/tiers, any text written on it)."
)


class Description(BaseModel):
    description: str = Field(min_length=1, max_length=400)


def summarize_spreadsheet(content: bytes, mime: str, filename: str | None = None) -> str:
    """Deterministic one-liner: format, sheet, row count, header and the first 3 rows."""
    name = (filename or "").lower()
    if mime.startswith("text/csv") or name.endswith(".csv"):
        rows = [r for r in csv.reader(io.StringIO(content.decode("utf-8-sig", errors="replace"))) if any(r)]
        return _summary("CSV", None, rows)
    if name.endswith(".xls") and not name.endswith(".xlsx"):
        return "Excel (old .xls format). Ask for .xlsx or CSV to read it"
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # corrupt / password-protected
        log.info("spreadsheet_unreadable", error=type(exc).__name__)
        return "Excel file (could not be read)"
    ws = wb.worksheets[0]
    rows = [
        ["" if v is None else str(v).strip() for v in row]
        for row in ws.iter_rows(values_only=True)
        if any(v not in (None, "") for v in row)
    ]
    extra = f" (+{len(wb.worksheets) - 1} more sheets)" if len(wb.worksheets) > 1 else ""
    wb.close()
    return _summary("Excel", ws.title + extra, rows)


def _summary(fmt: str, sheet: str | None, rows: list[list[str]]) -> str:
    where = f' sheet "{sheet}"' if sheet else ""
    if not rows:
        return f"{fmt}{where}: empty"
    header, body = [c for c in rows[0] if c], rows[1:]
    sample = "; ".join(" ".join(c for c in r if c)[:60] for r in body[:3])
    return f"{fmt}{where}: {len(body)} rows: {', '.join(header)[:120]}" + (f"; {sample}" if sample else "")


async def describe(model_gateway: ModelGateway, business_id: UUID, media_object_id: UUID, *, owner_file: bool) -> str:
    found = await gateway.media_path(media_object_id)
    if found is None:
        raise LookupError(f"media {media_object_id} not found")
    path, mime = found
    kind = gateway.file_kind(mime)
    content = await asyncio.to_thread(Path(path).read_bytes)
    if kind == "spreadsheet":
        filename = (
            next(
                (
                    f.filename
                    for f in await owner.search_files(business_id, limit=50)
                    if f.media_object_id == media_object_id
                ),
                None,
            )
            if owner_file
            else None
        )
        text_ = await asyncio.to_thread(summarize_spreadsheet, content, mime, filename)
    elif kind in ("image", "pdf") and len(content) <= MAX_INLINE_BYTES:
        try:
            result = await model_gateway.vision(
                "vision",
                OWNER_PROMPT if owner_file else CUSTOMER_PROMPT,
                [MediaInput(mime=mime, data=content)],
                Description,
                ctx=CallCtx(business_id=business_id, purpose="media.describe"),
            )
            text_ = result.description.strip()
        except ModelError as exc:
            log.warning("media_describe_failed", code=exc.code, kind=kind)
            raise  # the job retries (transient model errors); a dead job leaves "not described yet"
    else:
        text_ = {"pdf": "PDF (too large to read)", "image": "Image (too large to read)"}.get(kind, f"{kind} file")
    await gateway.set_media_description(business_id, media_object_id, text_)
    if owner_file:
        await owner.set_description(business_id, media_object_id, text_)
    await emit("media_describe", PipelineEventStatus.OK, business_id=business_id, kind=kind, owner=owner_file)
    return text_
